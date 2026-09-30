"""知识库：文档解析 → 切片 → 全文索引（SQLite FTS5）→ 检索。

- 关键词检索：中日韩文字按二元组（bigram）切分，英文/数字按单词小写；存入 kb_fts.tokens，BM25 排序。
  对型号、编号、条款号等精确匹配效果好，始终可用。
- 向量检索（配置 EMBEDDING_API_KEY 后启用）：上传后先建关键词索引立即可用，向量由后台任务补算；
  两路结果用 RRF 融合，再可选用 reranker 重排。向量/重排服务故障时自动退回关键词检索。
- 解析：纯文本类在本进程解码；PDF / Word / Excel / PPT 属于不可信的复杂格式，送到无网络沙箱 runner 中解析，
  避免解析器漏洞影响持有密钥的 agent-api。
- 权限：owner 可管理；visibility=public（仅管理员可设置）时所有访客可在对话中使用。
"""
import asyncio
import base64
import html
import logging
import re
import shutil
import time
import unicodedata
import uuid
from pathlib import PurePosixPath
from typing import Optional

import httpx
import numpy as np

from . import config, db, embeddings
from .skills import SkillError

logger = logging.getLogger("agent.kb")

KID_RE = re.compile(r"^[a-f0-9]{32}$")
TEXT_EXTS = {".txt", ".md", ".markdown", ".csv", ".json", ".html", ".htm", ".xml", ".yaml", ".yml", ".log"}
BINARY_EXTS = {".pdf", ".docx", ".xlsx", ".pptx"}
ALLOWED_EXTS = TEXT_EXTS | BINARY_EXTS

# ---------------- 分词 ----------------
_CJK = r"\u3040-\u30ff\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff\uac00-\ud7af"
_TOKEN_RE = re.compile(rf"[{_CJK}]+|[a-z0-9]+")


def _norm(text: str) -> str:
    return unicodedata.normalize("NFKC", text).lower()


def tokenize(text: str) -> list[str]:
    out = []
    for m in _TOKEN_RE.finditer(_norm(text)):
        w = m.group()
        if w[0].isascii():
            out.append(w)
        elif len(w) == 1:
            out.append(w)
        else:
            out.extend(w[i:i + 2] for i in range(len(w) - 1))
    return out


def build_match(query: str, max_terms: int = 40) -> str:
    """把自然语言/关键词查询转成 FTS5 表达式：各词项 OR，单个汉字用前缀匹配。"""
    seen, terms = set(), []
    for t in tokenize(query):
        if t in seen:
            continue
        seen.add(t)
        single_cjk = len(t) == 1 and not t.isascii()
        terms.append(f'"{t}"*' if single_cjk else f'"{t}"')
        if len(terms) >= max_terms:
            break
    return " OR ".join(terms)


# ---------------- 切片 ----------------
_SEPARATORS = [r"(?<=\n\n)", r"(?<=\n)", r"(?<=[。！？!?；;])", r"(?<= )"]


def _split(text: str, limit: int, overlap: int, level: int = 0) -> list[str]:
    if len(text) <= limit:
        return [text]
    if level >= len(_SEPARATORS):  # 实在切不开：定长窗口 + 重叠
        step = max(1, limit - overlap)
        return [text[i:i + limit] for i in range(0, len(text), step) if text[i:i + limit].strip()]
    parts = [p for p in re.split(_SEPARATORS[level], text) if p]
    if len(parts) == 1:
        return _split(text, limit, overlap, level + 1)
    out, buf = [], ""
    for p in parts:
        if len(p) > limit:
            if buf and len(buf) < limit // 4:   # 短标题等并入下一段，保留上下文
                p, buf = buf + p, ""
            if buf:
                out.append(buf)
                buf = ""
            out.extend(_split(p, limit, overlap, level + 1))
        elif len(buf) + len(p) > limit:
            out.append(buf)
            buf = p
        else:
            buf += p
    if buf:
        out.append(buf)
    return out


def chunk_text(text: str, limit: Optional[int] = None, overlap: Optional[int] = None) -> list[str]:
    limit = limit or config.KB_CHUNK_CHARS
    overlap = config.KB_CHUNK_OVERLAP if overlap is None else overlap
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\u3000]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return [c.strip() for c in _split(text, limit, overlap) if c.strip()]


# ---------------- 解析 ----------------
def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise SkillError("无法识别文本编码（请使用 UTF-8 或 GBK）")


def _html_to_text(s: str) -> str:
    s = re.sub(r"(?is)<(script|style|noscript)[^>]*>.*?</\1>", " ", s)
    s = re.sub(r"(?i)<br\s*/?>|</(p|div|li|h[1-6]|tr|section|article)>", "\n", s)
    s = re.sub(r"<[^>]+>", " ", s)
    s = html.unescape(s)
    return re.sub(r"[ \t]+", " ", s)


# 在沙箱中运行的解析脚本：读取工作目录中的输入文件，把纯文本写到 out.txt
_EXTRACT_CODE = r'''
import sys
src, ext = sys.argv[1], sys.argv[2]
out = []
if ext == ".pdf":
    from pypdf import PdfReader
    r = PdfReader(src)
    for i, p in enumerate(r.pages, 1):
        t = (p.extract_text() or "").strip()
        if t:
            out.append(f"[第 {i} 页]\n{t}")
elif ext == ".docx":
    import docx
    d = docx.Document(src)
    for block in d.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = "".join(n.text or "" for n in block.iter() if n.tag.endswith("}t")).strip()
            if t:
                out.append(t)
        elif tag == "tbl":
            rows = []
            for tr in block.iter():
                if tr.tag.endswith("}tr"):
                    cells = []
                    for tc in tr:
                        if tc.tag.endswith("}tc"):
                            cells.append("".join(n.text or "" for n in tc.iter() if n.tag.endswith("}t")).strip())
                    rows.append(" | ".join(cells))
            if rows:
                out.append("\n".join(rows))
elif ext == ".xlsx":
    import openpyxl
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    for ws in wb.worksheets:
        head, lines = None, []
        for row in ws.iter_rows(values_only=True):
            vals = ["" if v is None else str(v).strip() for v in row]
            if not any(vals):
                continue
            if head is None:
                head = vals
                continue
            pairs = [f"{head[i] if i < len(head) and head[i] else f'列{i+1}'}: {v}" for i, v in enumerate(vals) if v]
            lines.append("；".join(pairs))
        if head is not None:
            out.append(f"[工作表 {ws.title}] 列：" + "、".join(h for h in head if h) + ("\n" + "\n".join(lines) if lines else ""))
elif ext == ".pptx":
    from pptx import Presentation
    prs = Presentation(src)
    for i, s in enumerate(prs.slides, 1):
        texts = []
        for sh in s.shapes:
            if sh.has_text_frame:
                t = sh.text_frame.text.strip()
                if t:
                    texts.append(t)
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    texts.append(" | ".join(c.text.strip() for c in r.cells))
        if texts:
            out.append(f"[第 {i} 页幻灯片]\n" + "\n".join(texts))
open("out.txt", "w", encoding="utf-8").write("\n\n".join(out))
'''


async def _extract_in_sandbox(data: bytes, ext: str) -> str:
    src = "input" + ext
    payload = {
        "entry": {"type": "code", "code": _EXTRACT_CODE},
        "args": [src, ext],
        "stdin": "",
        "timeout": config.RUN_TIMEOUT,
        "work": {src: base64.b64encode(data).decode()},
        "skill": {},
    }
    try:
        async with httpx.AsyncClient(timeout=config.RUN_TIMEOUT + 30) as c:
            r = await c.post(f"{config.RUNNER_URL}/run", json=payload, headers={"X-Runner-Token": config.RUNNER_TOKEN})
    except httpx.HTTPError as e:
        logger.error("runner unreachable for kb extract: %r", e)
        raise SkillError("文档解析服务不可用，请稍后重试", 503)
    if r.status_code != 200:
        raise SkillError("文档解析服务繁忙，请稍后重试" if r.status_code == 503 else "文档解析失败", 503)
    res = r.json()
    out = (res.get("files") or {}).get("out.txt")
    if res.get("exit_code") != 0 or out is None:
        logger.info("kb extract failed: %s", (res.get("stderr") or "")[-500:])
        msg = "解析超时" if res.get("timed_out") else "文件解析失败，可能已损坏或加密"
        raise SkillError(msg, 422)
    return base64.b64decode(out).decode("utf-8", "replace")


async def extract_text(data: bytes, ext: str) -> str:
    if ext in BINARY_EXTS:
        return await _extract_in_sandbox(data, ext)
    text = _decode(data)
    if ext in (".html", ".htm", ".xml"):
        text = _html_to_text(text)
    return text


# ---------------- 权限 ----------------
def _check_kid(kid: str) -> None:
    if not KID_RE.match(kid or ""):
        raise SkillError("知识库不存在", 404)


def get_visible(vid: str, kid: str) -> dict:
    _check_kid(kid)
    kb = db.kb_get(kid)
    if not kb or (kb["owner"] != vid and kb["visibility"] != "public"):
        raise SkillError("知识库不存在", 404)
    return kb


def get_owned(vid: str, kid: str) -> dict:
    kb = get_visible(vid, kid)
    if kb["owner"] != vid:
        raise SkillError("公共知识库只读，不能修改", 403)
    return kb


def view(kb: dict, vid: Optional[str] = None) -> dict:
    out = {k: kb[k] for k in ("id", "name", "description", "visibility", "docs", "chars", "created_at", "updated_at")}
    if vid is not None:
        out["editable"] = kb["owner"] == vid
    return out


def resolve_selection(vid: str, kb_ids: list[str]) -> list[str]:
    """过滤出当前访客可用的知识库（去重、保序、限量）。不可见的直接丢弃。"""
    out = []
    for k in kb_ids or []:
        if not isinstance(k, str) or k in out or not KID_RE.match(k):
            continue
        kb = db.kb_get(k)
        if kb and (kb["owner"] == vid or kb["visibility"] == "public"):
            out.append(k)
        if len(out) >= config.KB_MAX_PER_SESSION:
            break
    return out


# ---------------- 管理 ----------------
def create(vid: str, name: str, description: str) -> dict:
    name = name.strip()
    if not name:
        raise SkillError("名称不能为空")
    if db.kb_count_owned(vid) >= config.KB_MAX_PER_VISITOR:
        raise SkillError(f"每位访客最多创建 {config.KB_MAX_PER_VISITOR} 个知识库", 403)
    return db.kb_create(vid, name[:60], description.strip()[:500])


def remove(kid: str) -> None:
    _check_kid(kid)
    db.kb_delete(kid)
    _vec_cache.pop(kid, None)
    shutil.rmtree(config.KB_DIR / kid, ignore_errors=True)


def safe_filename(name: str) -> str:
    name = PurePosixPath((name or "").replace("\\", "/")).name
    name = re.sub(r"[\x00-\x1f/\\:*?\"<>|]", "_", name).strip().lstrip(".")
    return name[:120] or "document.txt"


async def add_document(vid: str, kb: dict, filename: str, data: bytes) -> dict:
    """解析并入库一个文档，返回文档记录（附 truncated 标记）。"""
    filename = safe_filename(filename)
    ext = PurePosixPath(filename).suffix.lower()
    if ext not in ALLOWED_EXTS:
        raise SkillError("不支持的文件类型。支持：" + " ".join(sorted(ALLOWED_EXTS)))
    if not data:
        raise SkillError("文件为空")
    if len(db.kb_docs(kb["id"])) >= config.KB_MAX_DOCS:
        raise SkillError(f"每个知识库最多 {config.KB_MAX_DOCS} 篇文档", 403)
    if db.kb_owner_bytes(kb["owner"]) + len(data) > config.KB_VISITOR_QUOTA_MB * 1024 * 1024:
        raise SkillError(f"知识库空间不足（每位访客上限 {config.KB_VISITOR_QUOTA_MB} MB）", 413)

    text = (await extract_text(data, ext)).replace("\x00", "")
    truncated = len(text) > config.KB_DOC_MAX_CHARS
    if truncated:
        text = text[: config.KB_DOC_MAX_CHARS]
    chunks = chunk_text(text)
    if not chunks:
        raise SkillError("未能从文件中提取到文字（扫描版 PDF / 图片暂不支持）", 422)

    rows = [(c, " ".join(tokenize(c))) for c in chunks]
    doc_id = uuid.uuid4().hex
    d = config.KB_DIR / kb["id"]
    d.mkdir(parents=True, exist_ok=True)
    path = d / f"{doc_id}{ext}"
    path.write_bytes(data)
    try:
        doc = db.kb_doc_add(kb["id"], filename, ext, len(data), sum(len(c) for c in chunks), rows, doc_id=doc_id)
    except Exception:
        path.unlink(missing_ok=True)
        raise
    kick()
    return {**doc, "truncated": truncated}


def delete_document(kb: dict, doc_id: str) -> None:
    doc = db.kb_doc_get(kb["id"], doc_id) if KID_RE.match(doc_id or "") else None
    if not doc:
        raise SkillError("文档不存在", 404)
    db.kb_doc_delete(kb["id"], doc_id)
    (config.KB_DIR / kb["id"] / f"{doc_id}{doc['ext']}").unlink(missing_ok=True)


def document_file(kb: dict, doc_id: str) -> tuple[dict, bytes]:
    doc = db.kb_doc_get(kb["id"], doc_id) if KID_RE.match(doc_id or "") else None
    p = config.KB_DIR / kb["id"] / f"{doc_id}{doc['ext']}" if doc else None
    if not doc or not p.is_file():
        raise SkillError("文档不存在", 404)
    return doc, p.read_bytes()


# ---------------- 向量化（后台） ----------------
_kick = asyncio.Event()
_vec_cache: dict[str, tuple[tuple, np.ndarray, np.ndarray]] = {}  # kb_id -> (签名, chunk_ids, 矩阵)
_state = {"last_error": None, "last_error_at": 0.0}


def kick() -> None:
    """有新片段需要向量化时调用，唤醒后台任务。"""
    _kick.set()


async def embed_pending_once() -> int:
    """处理一批待向量化片段，返回处理数量。失败抛 EmbeddingError。"""
    if not embeddings.enabled():
        return 0
    model = embeddings.model_id()
    rows = db.kb_pending_chunks(model, config.EMBEDDING_BATCH)
    if not rows:
        return 0
    vecs = await embeddings.embed([r["content"] for r in rows])
    db.kb_set_embeddings([(r["id"], embeddings.to_blob(v)) for r, v in zip(rows, vecs)], model)
    return len(rows)


async def vector_loop() -> None:
    """后台常驻：有新文档或换模型时补算向量；出错指数退避，不影响关键词检索。"""
    backoff = 0.0
    while True:
        try:
            await asyncio.wait_for(_kick.wait(), timeout=300 if not backoff else backoff)
        except asyncio.TimeoutError:
            pass
        _kick.clear()
        if not embeddings.enabled():
            continue
        try:
            while await embed_pending_once():
                await asyncio.sleep(0.2)  # 给其他请求让出时间，也避免触发供应商限流
            backoff = 0.0
            _state["last_error"] = None
        except embeddings.EmbeddingError as e:
            backoff = min(600.0, (backoff or 15.0) * 2)
            _state.update(last_error=str(e)[:200], last_error_at=time.time())
            logger.warning("embedding failed, retry in %.0fs: %s", backoff, e)
        except Exception:
            backoff = min(600.0, (backoff or 15.0) * 2)
            logger.exception("vector loop error")


def vector_status() -> dict:
    out = embeddings.info()
    if out["enabled"]:
        out["pending"] = db.kb_pending_count(embeddings.model_id())
        out["error"] = _state["last_error"]
    return out


def _kb_matrix(kid: str, model: str) -> tuple[np.ndarray, np.ndarray]:
    sig = db.kb_vector_sig(kid, model) + (model,)
    hit = _vec_cache.get(kid)
    if hit and hit[0] == sig:
        return hit[1], hit[2]
    rows = db.kb_vectors(kid, model)
    ids = np.asarray([r["id"] for r in rows], dtype=np.int64)
    mat = embeddings.from_blobs([r["embedding"] for r in rows]) if rows else np.zeros((0, 0), np.float32)
    _vec_cache[kid] = (sig, ids, mat)
    if len(_vec_cache) > 200:  # 简单限量，防止常驻内存无限增长
        _vec_cache.pop(next(iter(_vec_cache)))
    return ids, mat


def _vector_recall(kb_ids: list[str], qvec: np.ndarray, limit: int) -> list[int]:
    model = embeddings.model_id()
    all_ids, all_scores = [], []
    for kid in kb_ids:
        ids, mat = _kb_matrix(kid, model)
        if len(ids) and mat.shape[1] == qvec.shape[0]:
            all_ids.append(ids)
            all_scores.append(mat @ qvec)
    if not all_ids:
        return []
    ids, scores = np.concatenate(all_ids), np.concatenate(all_scores)
    top = np.argsort(-scores)[:limit]
    return [int(ids[i]) for i in top]


def rrf(rankings: list[list[int]], k: int = 60) -> list[int]:
    """Reciprocal Rank Fusion：score = Σ 1/(k + rank)。"""
    score: dict[int, float] = {}
    for ranking in rankings:
        for rank, cid in enumerate(ranking):
            score[cid] = score.get(cid, 0.0) + 1.0 / (k + rank + 1)
    return sorted(score, key=lambda c: -score[c])


# ---------------- 检索 ----------------
async def search(kb_ids: list[str], query: str, top_k: int = 6) -> tuple[list[dict], str]:
    """混合检索：关键词(BM25) + 向量 → RRF 融合 → 可选重排。返回 (hits, 检索方式说明)。
    向量 / 重排任一环节失败都自动降级，不抛异常。"""
    top_k = max(1, min(int(top_k or 6), 12))
    recall = max(top_k, config.KB_RECALL)
    match = build_match(query)
    kw = [h["id"] for h in db.kb_search(kb_ids, match, recall)] if match else []
    rankings, modes = [kw], ["关键词"]

    if embeddings.enabled() and kb_ids:
        try:
            qvec = (await embeddings.embed([query], query=True))[0]
            vec = _vector_recall(kb_ids, qvec, recall)
            if vec:
                rankings.append(vec)
                modes.append("向量")
        except embeddings.EmbeddingError as e:
            logger.warning("vector search degraded to keyword: %s", e)

    fused = rrf(rankings) if len(rankings) > 1 else kw
    if not fused:
        return [], "+".join(modes)
    candidates = fused[:recall]
    rows = db.kb_chunks_by_ids(candidates)
    candidates = [c for c in candidates if c in rows]

    order = candidates[:top_k]
    if embeddings.rerank_enabled() and len(candidates) > 1:
        try:
            ranked = await embeddings.rerank(query, [rows[c]["content"] for c in candidates], top_k)
            if ranked:
                order = [candidates[i] for i, _ in ranked[:top_k]]
                modes.append("重排")
        except embeddings.EmbeddingError as e:
            logger.warning("rerank degraded: %s", e)

    names: dict[str, str] = {}
    hits = []
    for cid in order:
        h = dict(rows[cid])
        if h["kb_id"] not in names:
            kb = db.kb_get(h["kb_id"])
            names[h["kb_id"]] = kb["name"] if kb else ""
        h["kb_name"] = names[h["kb_id"]]
        hits.append(h)
    return hits, "+".join(modes)


def read_chunks(kb_ids: list[str], doc_id: str, start: int = 0, count: int = 3) -> tuple[dict, list[dict]]:
    doc = db.kb_doc_find(doc_id, kb_ids) if KID_RE.match(doc_id or "") else None
    if not doc:
        raise SkillError("文档不存在或不在当前会话选择的知识库中", 404)
    start = max(0, int(start or 0))
    count = max(1, min(int(count or 3), 8))
    return doc, db.kb_chunks_range(doc_id, start, count)

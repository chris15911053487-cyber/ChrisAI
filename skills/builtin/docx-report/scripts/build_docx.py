"""根据 JSON 规格生成 Word 文档。

用法：
  python build_docx.py 输出.docx            # 从 stdin 读取 JSON
  python build_docx.py 输出.docx --spec spec.json
"""
import argparse
import datetime
import json
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

FONT = "Microsoft YaHei"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)


def set_font(run, size=None, bold=None, color=None, name=FONT):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_rich(par, text, size=11):
    """支持 **粗体** 与 *斜体* 的简单行内格式。"""
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", str(text)):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2]); set_font(r, size, bold=True)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = par.add_run(part[1:-1]); set_font(r, size); r.italic = True
        else:
            r = par.add_run(part); set_font(r, size)


def shade(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def caption(doc, text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(str(text)); set_font(r, 9, color=RGBColor(0x59, 0x59, 0x59))


def add_toc(doc):
    p = doc.add_paragraph()
    r = p.add_run()
    for tag, text in (("begin", None), (None, 'TOC \\o "1-3" \\h \\z \\u'), ("separate", None), (None, None), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag); r._r.append(el)
        elif text:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text; r._r.append(el)
        else:
            t = OxmlElement("w:t"); t.text = "（在 Word 中右键此处 → 更新域，即可生成目录）"; r._r.append(t)


def build(spec, out):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2.4)

    normal = doc.styles["Normal"]
    normal.font.name = FONT; normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    for lvl, size in ((1, 16), (2, 14), (3, 12)):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = FONT; st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = ACCENT
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.paragraph_format.space_before = Pt(12); st.paragraph_format.space_after = Pt(6)

    # 标题区
    title = spec.get("title") or "未命名文档"
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(title), 22, bold=True, color=ACCENT)
    if spec.get("subtitle"):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(spec["subtitle"]), 13, color=RGBColor(0x40, 0x40, 0x40))
    info = " · ".join(x for x in [spec.get("author"), spec.get("date") or datetime.date.today().isoformat()] if x)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(info), 10, color=RGBColor(0x80, 0x80, 0x80))
    doc.core_properties.title = title
    if spec.get("author"):
        doc.core_properties.author = spec["author"]
    if spec.get("toc"):
        doc.add_heading("目录", level=1); add_toc(doc)
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    for i, b in enumerate(spec.get("blocks") or []):
        t = b.get("type")
        if t == "heading":
            doc.add_heading(str(b.get("text", "")), level=max(1, min(3, int(b.get("level", 1)))))
        elif t == "paragraph":
            add_rich(doc.add_paragraph(), b.get("text", ""))
        elif t in ("bullets", "numbered"):
            style = "List Bullet" if t == "bullets" else "List Number"
            for it in b.get("items") or []:
                add_rich(doc.add_paragraph(style=style), it)
        elif t == "table":
            headers = [str(h) for h in b.get("headers") or []]
            rows = [[str(c) for c in r] for r in b.get("rows") or []]
            ncol = max([len(headers)] + [len(r) for r in rows] + [1])
            table = doc.add_table(rows=0, cols=ncol)
            table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
            if headers:
                cells = table.add_row().cells
                for j in range(ncol):
                    cells[j].text = ""
                    set_font(cells[j].paragraphs[0].add_run(headers[j] if j < len(headers) else ""), 10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
                    shade(cells[j], "1F4E79")
            for k, r in enumerate(rows):
                cells = table.add_row().cells
                for j in range(ncol):
                    set_font(cells[j].paragraphs[0].add_run(r[j] if j < len(r) else ""), 10)
                    if k % 2:
                        shade(cells[j], "EAF1F8")
            if b.get("caption"):
                caption(doc, b["caption"])
            else:
                doc.add_paragraph()
        elif t == "quote":
            tbl = doc.add_table(rows=1, cols=1); c = tbl.rows[0].cells[0]; shade(c, "F2F6FA")
            add_rich(c.paragraphs[0], b.get("text", ""), 10.5)
            doc.add_paragraph()
        elif t == "image":
            path = str(b.get("path", ""))
            if not os.path.isfile(path):
                print(f"警告：第 {i + 1} 个块的图片不存在：{path}，已跳过", file=sys.stderr)
                continue
            doc.add_picture(path, width=Cm(float(b.get("width_cm", 14))))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if b.get("caption"):
                caption(doc, b["caption"])
        elif t == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        else:
            print(f"警告：未知块类型 {t!r}，已跳过", file=sys.stderr)

    # 页脚页码
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag); r._r.append(el)
        else:
            el = OxmlElement("w:instrText"); el.text = text; r._r.append(el)

    doc.save(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("output")
    ap.add_argument("--spec")
    a = ap.parse_args()
    out = os.path.basename(a.output) or "document.docx"
    if not out.lower().endswith(".docx"):
        out += ".docx"
    raw = open(a.spec, encoding="utf-8").read() if a.spec else sys.stdin.read()
    try:
        spec = json.loads(raw)
    except json.JSONDecodeError as e:
        sys.exit(f"JSON 规格解析失败：{e}")
    if not isinstance(spec, dict):
        sys.exit("JSON 规格必须是对象")
    build(spec, out)
    print(f"已生成 {out}（{os.path.getsize(out)} 字节，{len(spec.get('blocks') or [])} 个内容块）")


if __name__ == "__main__":
    main()

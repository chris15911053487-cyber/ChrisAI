"""根据 JSON 规格生成 Excel。

用法：
  python build_xlsx.py 输出.xlsx            # 从 stdin 读取 JSON
  python build_xlsx.py 输出.xlsx --spec data.json
"""
import argparse
import json
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
ALT_FILL = PatternFill("solid", fgColor="EAF1F8")
TOTAL_FILL = PatternFill("solid", fgColor="D9E2F3")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT = "Microsoft YaHei"


def width_of(v) -> float:
    s = str(v if v is not None else "")
    return sum(2 if ord(c) > 127 else 1 for c in s)


def safe_sheet_name(name, used):
    name = re.sub(r"[\[\]\*\?/\\:]", "_", str(name or "Sheet"))[:31] or "Sheet"
    base, n = name, 2
    while name in used:
        name = f"{base[:28]}_{n}"; n += 1
    used.add(name)
    return name


def build_sheet(ws, sh):
    headers = [str(h) for h in sh.get("headers") or []]
    rows = sh.get("rows") or []
    ncol = max([len(headers)] + [len(r) for r in rows] + [1])
    col_idx = {h: i + 1 for i, h in enumerate(headers)}

    if headers:
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(name=FONT, bold=True, color="FFFFFF"); c.fill = HEAD_FILL
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BORDER
        ws.freeze_panes = "A2"
    start = ws.max_row + 1 if headers else 1
    for k, r in enumerate(rows):
        ws.append(list(r) + [None] * (ncol - len(r)))
        for c in ws[ws.max_row]:
            c.font = Font(name=FONT); c.border = BORDER
            if k % 2:
                c.fill = ALT_FILL
    end = ws.max_row

    fmts = sh.get("number_formats") or {}
    for h, fmt in fmts.items():
        j = col_idx.get(h)
        if j:
            for i in range(start, end + 1):
                ws.cell(i, j).number_format = fmt

    tot = sh.get("totals")
    if tot and rows:
        tr = end + 2  # 空一行后写合计
        ws.cell(tr, 1, tot.get("label", "合计"))
        for h in tot.get("sum") or []:
            j = col_idx.get(h)
            if j:
                L = get_column_letter(j)
                ws.cell(tr, j, f"=SUM({L}{start}:{L}{end})").number_format = fmts.get(h, "General")
        for j in range(1, ncol + 1):
            c = ws.cell(tr, j); c.font = Font(name=FONT, bold=True); c.fill = TOTAL_FILL; c.border = BORDER

    if headers and rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(ncol)}{end}"

    for j in range(1, ncol + 1):
        vals = [headers[j - 1] if j <= len(headers) else ""] + [r[j - 1] if j <= len(r) else "" for r in rows[:500]]
        ws.column_dimensions[get_column_letter(j)].width = min(60, max(8, max(width_of(v) for v in vals) + 2))

    ch = sh.get("chart")
    if ch and headers and rows:
        kind = ch.get("type", "bar")
        chart = {"bar": BarChart, "line": LineChart, "pie": PieChart}.get(kind, BarChart)()
        chart.title = ch.get("title")
        chart.height, chart.width = 8, 16
        xj = col_idx.get(ch.get("x"), 1)
        ys = [col_idx[y] for y in ch.get("y") or [] if y in col_idx]
        if not ys:
            print("警告：chart.y 没有匹配的列，跳过图表", file=sys.stderr)
        else:
            if kind == "pie":
                ys = ys[:1]
            for j in ys:
                chart.add_data(Reference(ws, min_col=j, min_row=1, max_row=end), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=xj, min_row=start, max_row=end))
            ws.add_chart(chart, f"{get_column_letter(ncol + 2)}2")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("output")
    ap.add_argument("--spec")
    a = ap.parse_args()
    out = os.path.basename(a.output) or "workbook.xlsx"
    if not out.lower().endswith(".xlsx"):
        out += ".xlsx"
    raw = open(a.spec, encoding="utf-8").read() if a.spec else sys.stdin.read()
    try:
        spec = json.loads(raw)
    except json.JSONDecodeError as e:
        sys.exit(f"JSON 规格解析失败：{e}")
    sheets = spec.get("sheets") if isinstance(spec, dict) else None
    if not sheets:
        sys.exit("规格中缺少 sheets")
    wb = Workbook(); wb.remove(wb.active)
    used = set()
    for sh in sheets:
        build_sheet(wb.create_sheet(safe_sheet_name(sh.get("name"), used)), sh)
    wb.save(out)
    print(f"已生成 {out}：" + "，".join(f"{s.title}({s.max_row} 行)" for s in wb.worksheets))


if __name__ == "__main__":
    main()
